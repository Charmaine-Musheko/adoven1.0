from pathlib import Path
from textwrap import wrap

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image, ImageDraw, ImageFont


ROOT = Path("/Users/adoven/StudioProjects/adoven")
OUT_DIR = ROOT / "docs"
IMG_DIR = OUT_DIR / "diagrams"
DOCX = OUT_DIR / "Adoven Mobile App Technical Documentation.docx"


def ensure_dirs():
    OUT_DIR.mkdir(exist_ok=True)
    IMG_DIR.mkdir(exist_ok=True)


def font(size=26, bold=False):
    candidates = [
        "/System/Library/Fonts/Supplemental/Arial Bold.ttf" if bold else "/System/Library/Fonts/Supplemental/Arial.ttf",
        "/Library/Fonts/Arial Bold.ttf" if bold else "/Library/Fonts/Arial.ttf",
        "/System/Library/Fonts/Helvetica.ttc",
    ]
    for candidate in candidates:
        try:
            return ImageFont.truetype(candidate, size)
        except Exception:
            continue
    return ImageFont.load_default()


def draw_wrapped(draw, text, xy, max_width, fnt, fill=(20, 32, 28), line_gap=6):
    x, y = xy
    words = text.split()
    lines = []
    current = ""
    for word in words:
        trial = f"{current} {word}".strip()
        if draw.textbbox((0, 0), trial, font=fnt)[2] <= max_width:
            current = trial
        else:
            if current:
                lines.append(current)
            current = word
    if current:
        lines.append(current)
    for line in lines:
        draw.text((x, y), line, font=fnt, fill=fill)
        y += fnt.size + line_gap
    return y


def rounded_box(draw, xy, title, body, fill, outline=(64, 96, 78), width=2):
    x1, y1, x2, y2 = xy
    draw.rounded_rectangle(xy, radius=18, fill=fill, outline=outline, width=width)
    draw_wrapped(draw, title, (x1 + 18, y1 + 14), x2 - x1 - 36, font(22, True), fill=(18, 48, 38), line_gap=4)
    draw_wrapped(draw, body, (x1 + 18, y1 + 48), x2 - x1 - 36, font(17), fill=(42, 58, 51), line_gap=4)


def arrow(draw, start, end, fill=(42, 78, 60), width=4):
    draw.line([start, end], fill=fill, width=width)
    x1, y1 = start
    x2, y2 = end
    if abs(x2 - x1) >= abs(y2 - y1):
        direction = 1 if x2 > x1 else -1
        points = [(x2, y2), (x2 - 14 * direction, y2 - 8), (x2 - 14 * direction, y2 + 8)]
    else:
        direction = 1 if y2 > y1 else -1
        points = [(x2, y2), (x2 - 8, y2 - 14 * direction), (x2 + 8, y2 - 14 * direction)]
    draw.polygon(points, fill=fill)


def make_architecture_diagram():
    path = IMG_DIR / "system_architecture.png"
    img = Image.new("RGB", (1500, 900), "white")
    d = ImageDraw.Draw(img)
    d.text((50, 34), "Adoven System Architecture", font=font(34, True), fill=(16, 42, 34))

    boxes = {
        "user": (60, 140, 330, 270, "Customer", "Signs in, captures meters, receipts, readings, and unit balances."),
        "app": (430, 125, 760, 290, "Flutter Mobile App", "Material 3 UI, Riverpod state, Supabase client, REST API wrappers, receipt parser."),
        "auth": (880, 90, 1240, 220, "Supabase Auth", "Email/password sign-up, sign-in, session tokens, email confirmation, password reset deep links."),
        "db": (880, 285, 1240, 480, "Supabase Postgres", "profiles, properties, meters, readings, receipts, balance snapshots, usage snapshots, audit events, RLS policies."),
        "api": (430, 430, 760, 585, "DigitalOcean API", "Configurable API_BASE_URL with /api/analytics, /api/token, /api/report, /api/admin, /api/alert."),
        "bridge": (60, 570, 330, 725, "MQTT + Python Bridge", "Mosquitto receives smart meter telemetry; Python bridge writes readings to backend storage."),
        "meter": (60, 365, 330, 500, "Smart Meter", "Prototype meter sends energy readings and supports utility monitoring use cases."),
        "android": (880, 575, 1240, 720, "Android Runtime", "Internet/network permissions, launcher activity, and adoven:// auth/reset deep links."),
    }
    fills = {
        "user": (238, 247, 237), "app": (232, 241, 255), "auth": (244, 247, 236),
        "db": (250, 248, 240), "api": (244, 239, 255), "bridge": (255, 246, 232),
        "meter": (237, 246, 244), "android": (242, 247, 237)
    }
    for key, (x1, y1, x2, y2, title, body) in boxes.items():
        rounded_box(d, (x1, y1, x2, y2), title, body, fills[key])

    arrow(d, (330, 205), (430, 205))
    arrow(d, (760, 175), (880, 155))
    arrow(d, (760, 250), (880, 345))
    arrow(d, (760, 500), (880, 420))
    arrow(d, (330, 640), (430, 520))
    arrow(d, (195, 500), (195, 570))
    arrow(d, (1040, 575), (1040, 480))
    arrow(d, (595, 430), (595, 290))

    d.text((430, 315), "Active navigation: authenticated users enter the DashboardShell with Dashboard, Receipts, Analytics, and Settings tabs.", font=font(20), fill=(51, 65, 58))
    img.save(path)
    return path


def make_navigation_diagram():
    path = IMG_DIR / "navigation_flow.png"
    img = Image.new("RGB", (1500, 850), "white")
    d = ImageDraw.Draw(img)
    d.text((50, 34), "Application Navigation Flow", font=font(34, True), fill=(16, 42, 34))
    boxes = [
        ((80, 145, 330, 245), "Launch", "Preparing secure session"),
        ((430, 110, 760, 280), "Auth Landing", "Login, registration, password reset, password update"),
        ((870, 110, 1220, 280), "DashboardShell", "Bottom navigation wrapper for signed-in users"),
        ((90, 420, 350, 555), "Dashboard Tab", "Meter overview, add meter, add reading, update units"),
        ((430, 420, 690, 555), "Receipts Tab", "Search receipts, add receipt, parse raw receipt"),
        ((770, 420, 1030, 555), "Analytics Tab", "DigitalOcean analytics, energy overview, cost estimates"),
        ((1110, 420, 1370, 555), "Settings Tab", "Module notes, integration roadmap, sign out"),
    ]
    for xy, title, body in boxes:
        rounded_box(d, xy, title, body, (239, 247, 242))
    arrow(d, (330, 195), (430, 195))
    arrow(d, (760, 195), (870, 195))
    arrow(d, (1045, 280), (220, 420))
    arrow(d, (1045, 280), (560, 420))
    arrow(d, (1045, 280), (900, 420))
    arrow(d, (1045, 280), (1240, 420))
    img.save(path)
    return path


def make_erd_diagram():
    path = IMG_DIR / "entity_relationship_diagram.png"
    img = Image.new("RGB", (1600, 1050), "white")
    d = ImageDraw.Draw(img)
    d.text((50, 34), "Entity Relationship Diagram (Supabase/Postgres)", font=font(34, True), fill=(16, 42, 34))
    entities = {
        "auth.users": (70, 140, 350, 250, "id (PK)", "email, auth metadata"),
        "profiles": (470, 140, 750, 280, "id (PK/FK auth.users)", "full_name, created_at"),
        "properties": (870, 140, 1190, 300, "id (PK)", "owner_user_id (FK profiles), address_label, timezone"),
        "meters": (470, 400, 790, 590, "id (PK)", "property_id (FK properties), utility_type, meter_number, is_active"),
        "meter_readings": (70, 720, 390, 910, "id (PK)", "meter_id, reading_value, reading_timestamp, source, created_by"),
        "meter_balance_snapshots": (470, 720, 830, 930, "id (PK)", "meter_id, available_units, recorded_at, source, created_by"),
        "purchase_receipts": (910, 690, 1250, 945, "id (PK)", "meter_id, meter_number, token, receipt_number, amount, units, provider, idempotency_key, created_by"),
        "usage_snapshots": (1260, 400, 1550, 590, "id (PK)", "meter_id, period_start, period_end, units_used, estimated_cost"),
        "audit_events": (1270, 760, 1550, 930, "id (PK)", "entity_type, entity_id, action, actor_user_id, payload"),
    }
    for title, (x1, y1, x2, y2, pk, body) in entities.items():
        rounded_box(d, (x1, y1, x2, y2), title, f"{pk}\n{body}", (247, 250, 244))

    arrow(d, (350, 195), (470, 195))
    arrow(d, (750, 215), (870, 215))
    arrow(d, (1030, 300), (680, 400))
    arrow(d, (620, 590), (240, 720))
    arrow(d, (640, 590), (650, 720))
    arrow(d, (790, 510), (910, 790))
    arrow(d, (790, 490), (1260, 490))
    arrow(d, (1080, 945), (1410, 930))
    d.text((70, 975), "RLS rule pattern: authenticated users can read/write only rows owned by their Supabase user through profile, property, and meter ownership checks.", font=font(20), fill=(49, 64, 58))
    img.save(path)
    return path


def make_receipt_flow_diagram():
    path = IMG_DIR / "receipt_flow.png"
    img = Image.new("RGB", (1500, 760), "white")
    d = ImageDraw.Draw(img)
    d.text((50, 34), "Receipt Capture and Parsing Flow", font=font(34, True), fill=(16, 42, 34))
    boxes = [
        ((60, 165, 310, 300), "User", "Pastes raw prepaid receipt or enters fields manually."),
        ((390, 145, 700, 325), "AddReceiptSheet", "Select meter, parse raw payload, review token, receipt, amount, units, provider and optional VAT/address fields."),
        ((780, 145, 1080, 325), "ReceiptService", "Validates required values, rejects negative amounts/units, builds idempotency key."),
        ((1160, 145, 1440, 325), "Supabase Repository", "Writes purchase_receipts row for current authenticated user."),
        ((780, 470, 1080, 625), "Dashboard", "Uses purchases with usage and balance snapshots to estimate available units."),
    ]
    for xy, title, body in boxes:
        rounded_box(d, xy, title, body, (238, 247, 237))
    arrow(d, (310, 232), (390, 232))
    arrow(d, (700, 232), (780, 232))
    arrow(d, (1080, 232), (1160, 232))
    arrow(d, (1300, 325), (930, 470))
    img.save(path)
    return path


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_text(cell, text, bold=False):
    cell.text = ""
    p = cell.paragraphs[0]
    r = p.add_run(text)
    r.bold = bold
    r.font.size = Pt(9)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    for i, header in enumerate(headers):
        set_cell_text(hdr[i], header, bold=True)
        set_cell_shading(hdr[i], "DDEAD8")
        if widths:
            hdr[i].width = Inches(widths[i])
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            set_cell_text(cells[i], str(value))
            if widths:
                cells[i].width = Inches(widths[i])
    doc.add_paragraph()
    return table


def add_bullets(doc, items):
    for item in items:
        doc.add_paragraph(item, style="List Bullet")


def add_numbered(doc, items):
    for item in items:
        doc.add_paragraph(item, style="List Number")


def configure_doc(doc):
    section = doc.sections[0]
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)
    styles = doc.styles
    styles["Normal"].font.name = "Arial"
    styles["Normal"].font.size = Pt(10)
    for name, size, color in [
        ("Title", 24, "17301F"),
        ("Heading 1", 16, "17301F"),
        ("Heading 2", 13, "35693F"),
        ("Heading 3", 11, "17301F"),
    ]:
        style = styles[name]
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.font.bold = True


def build_document(diagrams):
    doc = Document()
    configure_doc(doc)

    title = doc.add_paragraph()
    title.style = doc.styles["Title"]
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("Adoven Mobile Application Technical Documentation")
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.add_run("From initial setup to current implementation state").italic = True
    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta.add_run("Prepared: 28 May 2026 | Project: /Users/adoven/StudioProjects/adoven")
    doc.add_paragraph()

    doc.add_heading("Purpose", level=1)
    doc.add_paragraph(
        "This document explains the Adoven mobile application from the start of the build through the current codebase. "
        "It covers what was created, how the Flutter app is structured, how it connects to Supabase and backend APIs, "
        "what data entities exist, what the active screens do, and what remains for future hardening."
    )

    doc.add_heading("Executive Summary", level=1)
    add_bullets(doc, [
        "Built a Flutter/Dart mobile app named Adoven with a Material 3 interface and Riverpod state management.",
        "Added Supabase authentication for email/password sign-up, sign-in, email confirmation, session handling, password reset, and password update.",
        "Created a Supabase/Postgres data model for profiles, properties, meters, meter readings, purchase receipts, balance snapshots, usage snapshots, and audit events.",
        "Protected Supabase data with row-level security policies so authenticated users can only access their own profile, property, meter, receipt, reading, and balance data.",
        "Implemented a dashboard for linked utility meters, available units, latest readings, and reading trends.",
        "Implemented receipt capture with raw prepaid receipt parsing, validation, masked token display, filtering, and Supabase persistence.",
        "Added a DigitalOcean/API integration layer for analytics and legacy endpoints such as token, report, admin, and alert.",
        "Configured Android permissions and deep links for internet access and Supabase auth/password reset redirects.",
        "Added unit/widget/integration tests around dashboard display, receipt parsing, receipt validation, receipt listing, and receipt creation flow.",
    ])

    doc.add_heading("Technology Stack", level=1)
    add_table(doc, ["Layer", "Technology", "Purpose"], [
        ("Mobile app", "Flutter + Dart", "Cross-platform UI and app logic."),
        ("State management", "flutter_riverpod", "Providers for dashboard, meters, receipts, filters, and repository overrides in tests."),
        ("Authentication", "Supabase Auth", "Email/password auth, session token handling, email confirmation, password reset."),
        ("Database", "Supabase Postgres", "Structured utility, receipt, meter, profile, and audit data."),
        ("Security", "RLS policies + Supabase sessions", "Restrict user data to the authenticated owner."),
        ("Backend API", "HTTP endpoints via API_BASE_URL", "Analytics and older token/report/admin/alert integrations."),
        ("IoT bridge", "Mosquitto MQTT + Python bridge", "Prototype path for meter telemetry ingestion into backend storage."),
        ("Android config", "AndroidManifest.xml", "Internet/network permissions and adoven:// deep links."),
        ("Testing", "flutter_test + integration_test", "Widget, service, parser, screen, and integration coverage."),
    ], widths=[1.5, 2.0, 4.0])

    doc.add_heading("Project Timeline", level=1)
    add_numbered(doc, [
        "Created the Flutter project skeleton with Android, iOS, web, macOS, Linux, and Windows directories.",
        "Added dependencies for HTTP, Supabase, Riverpod, secure storage/shared preferences, crypto, and decimal handling.",
        "Configured Supabase options through Dart defines so the app can point to different environments without hard-coding production settings.",
        "Built authentication screens and controller logic for sign-in, registration, email confirmation messaging, password reset, and password update.",
        "Created the Supabase database schema with utility-specific entities and row-level security.",
        "Built the authenticated app shell with Dashboard, Receipts, Analytics, and Settings tabs.",
        "Built the dashboard meter workflow: add a meter, capture available units, capture readings, and view trend narratives.",
        "Built receipt capture: paste raw receipt, parse fields, validate values, save receipt, list receipts, filter receipts, and mask token/meter values.",
        "Added analytics integration to a backend API hosted behind API_BASE_URL, including current power, kWh movement, estimates, opportunities, and split items.",
        "Added Android internet/network permissions and deep links for adoven://auth/callback and adoven://reset-password.",
        "Added tests for receipt parsing, receipt service validation, receipt screen display, add receipt sheet defaults, dashboard display, and receipt creation flow.",
    ])

    doc.add_heading("System Architecture", level=1)
    doc.add_paragraph("The active system combines the Flutter mobile client, Supabase auth/data services, a configurable backend API, and an IoT telemetry bridge.")
    doc.add_picture(str(diagrams["architecture"]), width=Inches(7.3))

    doc.add_heading("Application Navigation", level=1)
    doc.add_paragraph("The app decides between authentication and the signed-in dashboard shell through AppRouter. Once authenticated, users use a four-tab bottom navigation shell.")
    doc.add_picture(str(diagrams["navigation"]), width=Inches(7.3))

    doc.add_heading("Entity Relationship Diagram", level=1)
    doc.add_paragraph("The main persistent data model is in Supabase/Postgres. The schema uses ownership through auth.users -> profiles -> properties -> meters, then attaches readings, receipts, balances, usage, and audit records to that ownership chain.")
    doc.add_picture(str(diagrams["erd"]), width=Inches(7.3))

    doc.add_heading("Core Entities", level=2)
    add_table(doc, ["Entity", "Key Fields", "Role in the App"], [
        ("auth.users", "id, email, auth metadata", "Supabase-managed identity source."),
        ("profiles", "id, full_name, created_at", "Application profile linked one-to-one to auth.users."),
        ("properties", "owner_user_id, address_label, timezone", "Groups meters under a user-owned address/property."),
        ("meters", "property_id, utility_type, meter_number, is_active", "Represents electricity or water meters linked to a property."),
        ("meter_readings", "meter_id, reading_value, reading_timestamp, source, created_by", "Stores manual or imported readings used for trend calculations."),
        ("meter_balance_snapshots", "meter_id, available_units, recorded_at, source, created_by", "Stores available-unit snapshots for dashboard balances."),
        ("purchase_receipts", "meter_id, token, receipt_number, amount, units, provider, idempotency_key", "Stores prepaid purchase receipt details and raw receipt payloads."),
        ("usage_snapshots", "meter_id, period_start, period_end, units_used, estimated_cost", "Stores summarized usage/cost periods."),
        ("audit_events", "entity_type, entity_id, action, actor_user_id, payload", "Logs insert/update activity for receipts."),
    ], widths=[1.65, 2.8, 3.35])

    doc.add_heading("Receipt Flow", level=1)
    doc.add_picture(str(diagrams["receipt"]), width=Inches(7.3))
    doc.add_paragraph(
        "The receipt flow is one of the most complete feature slices. A user selects a linked meter, pastes a raw prepaid receipt, "
        "parses the payload into structured fields, reviews the details, and saves it to Supabase. The saved units feed dashboard calculations."
    )

    doc.add_heading("Backend and API Work", level=1)
    add_table(doc, ["API/Service", "Current Use", "Implementation Detail"], [
        ("Supabase Auth", "Registration, login, session checks, reset links", "ConsumerService wraps signUp, signInWithPassword, resetPasswordForEmail, updateUser, and signOut."),
        ("Supabase Postgres", "Meters, readings, receipts, balance snapshots", "Repositories use the Supabase Dart client and rely on RLS for ownership security."),
        ("DigitalOcean API", "Analytics and older endpoints", "API_BASE_URL defaults to http://164.92.194.34:7234/api and can be overridden with --dart-define."),
        ("/api/analytics", "Analytics tab", "AnalyticsRepository fetches devices, bills, opportunities, and split items."),
        ("/api/token", "Legacy token client", "TokenService can fetch/create Token DTOs."),
        ("/api/report", "Legacy report client", "ReportService can fetch/create Report DTOs."),
        ("/api/admin", "Legacy admin client", "AdminService can fetch/create Admin DTOs."),
        ("/api/alert", "Legacy alert client", "AlertService can fetch/create Alert DTOs."),
        ("MQTT/Python bridge", "Smart meter telemetry path", "Prototype bridge receives meter data from Mosquitto and writes readings to backend storage."),
    ], widths=[1.6, 2.35, 3.85])
    doc.add_paragraph(
        "Important distinction: the Flutter app does not contain all backend API server code. It contains client-side wrappers and repositories that call Supabase and configured backend endpoints."
    )

    doc.add_heading("Security Work", level=1)
    add_bullets(doc, [
        "Authentication state is centralized in AuthController and AppRouter uses it to keep unauthenticated users on the auth screens.",
        "Supabase sessions provide access tokens; AuthStorage reads the current Supabase session.",
        "Registration stores user metadata such as first name, last name, contact number, and address in Supabase Auth metadata.",
        "Password reset uses the adoven://reset-password deep link.",
        "AndroidManifest.xml registers adoven://auth/callback and adoven://reset-password intent filters.",
        "Supabase RLS policies enforce profile, property, meter, reading, receipt, balance, usage, and audit visibility by authenticated user ownership.",
        "Receipt logging redacts sensitive receipt payloads in debug output.",
        "The current main.dart accepts bad certificates through HttpOverrides. That is convenient for prototype testing, but it should be removed or restricted before production release.",
        "Android uses android:usesCleartextTraffic=true. That is helpful for local/prototype HTTP endpoints, but production should use HTTPS and disable cleartext traffic.",
    ])

    doc.add_heading("Screen Inventory", level=1)
    add_table(doc, ["Screen", "Status", "What Is On The Screen", "Main Actions/Data"], [
        ("Launch Screen", "Active", "Progress indicator and message: Preparing your secure session.", "Shown while AuthController initializes session state."),
        ("Auth Landing", "Active", "Adoven title, supporting text, animated auth form area.", "Switches between login, register, reset password, and update password views."),
        ("Login Screen", "Active", "Email field, password field with visibility toggle, Sign in button, Forgot password link, Create account link.", "Validates email/password and calls ConsumerService.loginConsumer through AuthController."),
        ("Register Screen", "Active", "First name, last name, email, password, contact number, address, Create account button.", "Creates Supabase auth user, handles email confirmation requirement."),
        ("Password Reset Request", "Active", "Email field, Send reset link button, Back to sign in link.", "Calls Supabase resetPasswordForEmail with adoven://reset-password redirect."),
        ("Password Update", "Active", "New password and confirm password fields.", "Updates Supabase user password after password recovery deep link."),
        ("Dashboard Shell", "Active", "Bottom navigation with Dashboard, Receipts, Analytics, Settings.", "Keeps tab state with IndexedStack."),
        ("Dashboard Tab", "Active", "Home app bar, refresh action, overview hero, meter summary cards, Add Meter FAB.", "Loads meter summaries, total units, latest readings, trends, add meter, add reading, update units."),
        ("Add Meter Sheet", "Active modal", "Address/property, meter type, meter number, optional current units, optional current reading.", "Creates property if needed, inserts meter, optional balance snapshot and reading."),
        ("Record Reading Sheet", "Active modal", "Single meter reading entry form.", "Inserts meter_readings row and refreshes dashboard."),
        ("Update Units Sheet", "Active modal", "Available units entry form.", "Inserts meter_balance_snapshots row and refreshes dashboard."),
        ("Receipts Tab", "Active", "Search field, receipt list, empty/error states, Add Receipt FAB.", "Lists receipts, filters by provider/meter/receipt number, masks sensitive token/meter display."),
        ("Add Receipt Sheet", "Active modal", "Meter selector, raw payload box, parse buttons, token, receipt number, amount, units, provider, VAT, customer, address, Save Receipt.", "Parses raw prepaid receipts and saves purchase_receipts."),
        ("Analytics Tab", "Active", "Refresh action, analytics cards, energy overview, cost bars, device monitoring, savings pipeline, shared balance health.", "Fetches /api/analytics from backend API."),
        ("Settings Tab", "Active", "Utility workspace summary, product modules, next integrations, Sign out button.", "Displays roadmap-style items and signs the user out."),
        ("HomeScreen", "Legacy/not routed", "Older meter list with refresh, sign out, Add Meter FAB.", "Uses MeterService REST/Supabase API style; not reached by current AppRouter."),
        ("AddMeterScreen", "Legacy/not routed", "Municipality, address, units, token ID, consumer ID.", "Creates older Meter DTO through MeterService."),
        ("MeterDetailsScreen", "Legacy/not routed", "Municipality, address, units, token ID, consumer ID.", "Read-only older meter detail page."),
        ("IncidentScreen", "Stub/not routed", "Dummy incident text and Back to Home button.", "Placeholder for future incident logging."),
        ("ProfileScreen", "Stub/not routed", "Dummy profile text and Back to Home button.", "Placeholder for future profile management."),
    ], widths=[1.45, 1.15, 3.0, 2.2])

    doc.add_heading("What Happens On The Main Screens", level=1)
    doc.add_heading("Authentication Screens", level=2)
    add_bullets(doc, [
        "The auth landing page uses a soft gradient background and shows the current auth form with AnimatedSwitcher.",
        "Login validates an email-like value and a non-empty password before calling Supabase sign-in.",
        "Registration collects personal contact details and requires a password of at least eight characters.",
        "Password reset sends an email through Supabase; update password handles the recovery deep-link path.",
    ])
    doc.add_heading("Dashboard Screen", level=2)
    add_bullets(doc, [
        "If no meters exist, the user sees a setup prompt explaining what can be added.",
        "If meters exist, the screen shows total linked meters and combined units on hand.",
        "Each meter card shows utility type, property label, meter number, available units, latest reading, and a trend label.",
        "Add Reading and Update Units open bottom sheets for meter-specific entry.",
    ])
    doc.add_heading("Receipts Screen", level=2)
    add_bullets(doc, [
        "Users can filter receipts by provider, meter number, or receipt number.",
        "Receipts show provider, masked token, masked meter number, units, amount, receipt number, and timestamp.",
        "The add receipt sheet can parse raw receipt text into structured fields before saving.",
    ])
    doc.add_heading("Analytics Screen", level=2)
    add_bullets(doc, [
        "The analytics tab reads JSON from /api/analytics.",
        "It shows current power, last 24-hour kWh, last 30-day kWh, estimated monthly cost, cost bars, device monitoring, savings opportunities, and split-balance items when returned by the API.",
    ])
    doc.add_heading("Settings Screen", level=2)
    add_bullets(doc, [
        "Settings currently acts as a roadmap and sign-out page.",
        "It lists product modules and planned integrations, then provides a Sign out button.",
    ])

    doc.add_heading("Testing Completed", level=1)
    add_table(doc, ["Test File", "Coverage"], [
        ("test/widget_test.dart", "Dashboard shows utility experience, meter card, overview copy, and trend label."),
        ("test/features/receipts/receipt_payload_parser_test.dart", "Parses raw prepaid receipt text into token, meter number, receipt number, provider, amount, units, VAT, address, and idempotency key."),
        ("test/features/receipts/receipt_service_test.dart", "Rejects negative values and verifies receipt parsing through the service layer."),
        ("test/features/receipts/receipts_screen_test.dart", "Displays saved receipt details and masked token data."),
        ("test/features/receipts/add_receipt_sheet_test.dart", "Shows add receipt guidance, sample receipt controls, and linked meter selection."),
        ("integration_test/receipt_creation_flow_test.dart", "Simulates parse-and-submit receipt flow with an in-memory repository."),
    ], widths=[2.75, 5.0])

    doc.add_heading("Current Limitations and Next Steps", level=1)
    add_bullets(doc, [
        "Replace prototype HTTP/certificate allowances with production TLS enforcement.",
        "Decide whether the legacy REST DTO screens should be removed, migrated, or connected into the active navigation.",
        "Promote incident logging and profile management from placeholders into active screens if they remain thesis requirements.",
        "Connect token submission to the real utility/SAP/MV90 integration path once an approved production API exists.",
        "Add role-based administration if administrator approval, user management, and incident management remain in scope.",
        "Add load testing and API monitoring for the DigitalOcean backend before a production deployment.",
        "Update the README from starter Flutter text to project-specific deployment and operations instructions.",
    ])

    doc.add_heading("Appendix: Key Local Files", level=1)
    add_table(doc, ["Area", "Files"], [
        ("App startup", "lib/main.dart, lib/app/app.dart, lib/app/router.dart"),
        ("Auth", "lib/ui/authentication/*, lib/api/consumer_api.dart, lib/api/auth_service.dart"),
        ("Dashboard", "lib/features/dashboard/presentation/screens/dashboard_screen.dart"),
        ("Receipts", "lib/features/receipts/**"),
        ("Utilities/meters", "lib/features/utilities/**"),
        ("Analytics", "lib/features/analytics/**"),
        ("Settings", "lib/features/settings/presentation/screens/settings_screen.dart"),
        ("API constants", "lib/api/api_constants/api_constants.dart"),
        ("Supabase schema", "supabase/migrations/*.sql"),
        ("Android configuration", "android/app/src/main/AndroidManifest.xml"),
        ("Tests", "test/**, integration_test/**"),
    ], widths=[2.0, 5.8])

    doc.save(DOCX)


def main():
    ensure_dirs()
    diagrams = {
        "architecture": make_architecture_diagram(),
        "navigation": make_navigation_diagram(),
        "erd": make_erd_diagram(),
        "receipt": make_receipt_flow_diagram(),
    }
    build_document(diagrams)
    print(DOCX)


if __name__ == "__main__":
    main()
